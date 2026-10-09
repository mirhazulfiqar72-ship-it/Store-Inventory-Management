"""Safe Update & Add import for Inventory Codes, v1.0.89.

Does not delete codes, modify MTO, or change GRN/Demand/Issue history.
The application's OnlineConnection commits changed rows through existing Firebase sync.
"""
from __future__ import annotations
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import csv
import math
import os
import re
import sqlite3

@dataclass(frozen=True)
class Item:
    code: str
    description: str
    uom: str | None
    opening: float | None
    location: str

ALIASES = {
    "code": {"itemcode", "code", "inventorycode", "stockcode", "materialcode", "productcode"},
    "description": {"description", "itemdescription", "itemname", "name", "materialdescription"},
    "uom": {"uom", "unit", "unitofmeasure", "units"},
    "opening": {"openingbalance", "openingbal", "openingqty", "openingquantity", "opening", "openingstock"},
}

def _clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\u00a0", " ").strip()

def _header(value):
    return re.sub(r"[^a-z0-9]", "", _clean(value).casefold())

def _fields(values):
    result = {}
    for i, value in enumerate(values):
        for name, possibilities in ALIASES.items():
            if _header(value) in possibilities and name not in result:
                result[name] = i
    return result if {"code", "description"} <= result.keys() else None

def _get(row, idx):
    return row[idx] if idx is not None and idx < len(row) else None

def _parse(rows):
    items, seen = [], set()
    current_group, mapping = None, None
    for location, row in rows:
        group = location.split("!")[0]
        if group != current_group:
            current_group, mapping = group, None
        if not row or not any(_clean(v) for v in row):
            continue
        new_mapping = _fields(row)
        if new_mapping:
            mapping = new_mapping
            continue
        if mapping is None:
            continue
        code = _clean(_get(row, mapping["code"]))
        description = _clean(_get(row, mapping["description"]))
        if not code and not description:
            continue
        if not code or not description:
            raise ValueError(f"{location}: Item Code and Description are required")
        if not re.fullmatch(r"[\w./ -]{3,45}", code) or len(description) > 400:
            raise ValueError(f"{location}: invalid code or description")
        if code.casefold() in seen:
            raise ValueError(f"{location}: duplicate Item Code {code}")
        seen.add(code.casefold())
        uom = _clean(_get(row, mapping.get("uom"))) if "uom" in mapping else ""
        opening_raw = _clean(_get(row, mapping.get("opening"))) if "opening" in mapping else ""
        opening = None
        if opening_raw:
            try:
                opening = float(opening_raw.replace(",", ""))
            except ValueError as error:
                raise ValueError(f"{location}: invalid Opening Balance {opening_raw!r}") from error
            if not math.isfinite(opening) or opening < 0:
                raise ValueError(f"{location}: Opening Balance must be non-negative")
        items.append(Item(code, description, uom or None, opening, location))
    if not items:
        raise ValueError("No rows read. Required headings: Item Code and Description")
    return items

def read_file(path):
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            def rows():
                for sheet in book.worksheets:
                    for index, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                        yield f"{sheet.title}!{index}", cells
            return _parse(rows())
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path)
        def rows():
            for sheet in book.sheets():
                for index in range(sheet.nrows):
                    yield f"{sheet.name}!{index+1}", sheet.row_values(index)
        return _parse(rows())
    if suffix == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as file:
            return _parse((f"CSV!{i}", cells) for i, cells in enumerate(csv.reader(file), 1))
    if suffix == ".pdf":
        import fitz
        doc = fitz.open(path)
        try:
            def rows():
                for page_no, page in enumerate(doc, 1):
                    found_table = False
                    try:
                        for table_no, table in enumerate(page.find_tables().tables, 1):
                            found_table = True
                            for index, row in enumerate(table.extract(), 1):
                                yield f"Page{page_no}-Table{table_no}!{index}", row
                    except (AttributeError, ValueError):
                        pass
                    if found_table:
                        continue
                    for index, line in enumerate((page.get_text("text") or "").splitlines(), 1):
                        fields = re.split(r"\t+| {2,}", line.strip())
                        if len(fields) > 1:
                            yield f"Page{page_no}!{index}", fields
            try:
                return _parse(rows())
            except ValueError as error:
                if "No rows read" in str(error):
                    raise ValueError("This PDF has no readable inventory table. Please use a selectable-text/table PDF or Excel; scanned PDFs are not supported.") from error
                raise
        finally:
            doc.close()
    raise ValueError("Select an Excel (.xlsx/.xls), CSV or selectable-text PDF file")

def preview(conn, items):
    rows = {str(r[0]).casefold(): r for r in conn.execute(
        "SELECT code,description,uom,opening_qty,item_type FROM items")}
    try:
        mto_codes = {str(r[0]).casefold() for r in conn.execute("SELECT code FROM mto_items")}
    except sqlite3.OperationalError:
        mto_codes = set()
    output = []
    for item in items:
        existing = rows.get(item.code.casefold())
        if item.code.casefold() in mto_codes or (existing and str(existing[4] or "").upper() == "MTO"):
            action = "MTO conflict"
        elif existing is None:
            action = "New"
        elif ((existing[1] or "") == item.description
              and (item.uom is None or (existing[2] or "") == item.uom)
              and (item.opening is None or abs(float(existing[3] or 0) - item.opening) < 1e-10)):
            action = "Unchanged"
        else:
            action = "Update"
        output.append((item, action, existing[0] if existing else item.code))
    return output

def _backup(conn, backup_dir):
    os.makedirs(backup_dir, exist_ok=True)
    path = os.path.join(backup_dir, "inventory_before_import_"
                        + datetime.now().strftime("%Y%m%d_%H%M%S_%f") + ".db")
    destination = sqlite3.connect(path)
    try:
        getattr(conn, "_conn", conn).backup(destination)
    finally:
        destination.close()
    return path

def apply(conn, imports, backup_dir):
    # Re-evaluate status, so a stale Preview cannot accidentally overwrite MTO.
    changes = preview(conn, imports)
    if not changes:
        raise ValueError("No rows to import")
    if any(action == "MTO conflict" for _, action, _ in changes):
        raise ValueError("Uploaded file contains MTO codes. Remove those rows and try again.")
    backup = _backup(conn, backup_dir)  # Complete backup BEFORE ANY data changes
    stats = {"new": 0, "updated": 0, "unchanged": 0, "backup": backup}
    try:
        for item, action, database_code in changes:
            if action == "Unchanged":
                stats["unchanged"] += 1
                continue
            if action == "Update":
                old = conn.execute("SELECT uom,opening_qty FROM items WHERE code=?", (database_code,)).fetchone()
                conn.execute(
                    "UPDATE items SET description=?,uom=?,opening_qty=? WHERE code=?",
                    (item.description, item.uom if item.uom is not None else old[0],
                     item.opening if item.opening is not None else old[1], database_code))
                stats["updated"] += 1
            else:
                conn.execute(
                    "INSERT INTO items(code,description,uom,category,opening_qty,min_level,item_type,mto_opening_qty) "
                    "VALUES(?,?,?,?,?,?,?,?)",
                    (item.code, item.description, item.uom or "Number", "",
                     item.opening if item.opening is not None else 0, 0, "Local", 0))
                stats["new"] += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    stats["online_pending"] = bool(getattr(getattr(conn, "sync", None), "pending_error", None))
    return stats
